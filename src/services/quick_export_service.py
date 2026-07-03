"""
Quick report export — Word/Excel/PDF versions of the current dataset plus a
summary statistics section, for the Reports menu's one-click export actions.
Deliberately lightweight compared to services/report_service.py's full
section-based HTML report builder: these are a fast "just get me a file"
path, not a configurable report.
"""
from pathlib import Path

import polars as pl

NUMERIC_DTYPES = (
    pl.Float32, pl.Float64, pl.Int8, pl.Int16, pl.Int32, pl.Int64,
    pl.UInt8, pl.UInt16, pl.UInt32, pl.UInt64,
)
PREVIEW_ROWS = 50


def _numeric_columns(df: pl.DataFrame):
    return [c for c in df.columns if df[c].dtype in NUMERIC_DTYPES]


def _summary_rows(df: pl.DataFrame):
    rows = []
    for col in _numeric_columns(df):
        s = df[col].drop_nulls()
        if s.is_empty():
            continue
        std = s.std()
        rows.append([
            col, str(s.len()), f"{s.mean():.4f}", f"{std:.4f}" if std is not None else "N/A",
            f"{s.min():.4f}", f"{s.max():.4f}",
        ])
    return rows


def export_word(df: pl.DataFrame, dataset_name: str, output_path: str) -> str:
    """Export dataset name, summary statistics, and a row preview to .docx."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(f"Data Report: {dataset_name}", level=1)
    doc.add_paragraph(f"{df.height:,} rows x {df.width} columns")

    doc.add_heading("Summary Statistics", level=2)
    summary_rows = _summary_rows(df)
    if summary_rows:
        headers = ["Column", "Count", "Mean", "Std Dev", "Min", "Max"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
        for row in summary_rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
    else:
        doc.add_paragraph("No numeric columns to summarize.")

    doc.add_heading(f"Data Preview (first {PREVIEW_ROWS} rows)", level=2)
    preview = df.head(PREVIEW_ROWS)
    table = doc.add_table(rows=1, cols=preview.width)
    table.style = "Light Grid Accent 1"
    for i, col in enumerate(preview.columns):
        table.rows[0].cells[i].text = col
    for row in preview.iter_rows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
        for cell in cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

    doc.save(output_path)
    return output_path


def export_excel(df: pl.DataFrame, dataset_name: str, output_path: str) -> str:
    """Export the full dataset to .xlsx, with a second sheet of summary stats."""
    pdf = df.to_pandas()
    import pandas as pd

    summary_rows = _summary_rows(df)
    summary_df = pd.DataFrame(summary_rows, columns=["Column", "Count", "Mean", "Std Dev", "Min", "Max"]) \
        if summary_rows else pd.DataFrame()

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pdf.to_excel(writer, sheet_name=dataset_name[:31] or "Data", index=False)
        if not summary_df.empty:
            summary_df.to_excel(writer, sheet_name="Summary", index=False)

    return output_path


def export_pdf(df: pl.DataFrame, dataset_name: str, output_path: str) -> str:
    """Export dataset name, summary statistics, and a row preview to .pdf
    using matplotlib's PDF backend — no extra PDF library dependency needed
    since matplotlib is already used elsewhere in this app."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    with PdfPages(output_path) as pdf_pages:
        # Title + summary page
        fig, ax = plt.subplots(figsize=(11, 8.5))
        ax.axis("off")
        ax.text(0.5, 0.95, f"Data Report: {dataset_name}", ha="center", fontsize=16, weight="bold")
        ax.text(0.5, 0.90, f"{df.height:,} rows x {df.width} columns", ha="center", fontsize=10, color="gray")

        summary_rows = _summary_rows(df)
        if summary_rows:
            headers = ["Column", "Count", "Mean", "Std Dev", "Min", "Max"]
            table = ax.table(cellText=summary_rows, colLabels=headers, loc="center", bbox=[0.05, 0.15, 0.9, 0.65])
            table.auto_set_font_size(False)
            table.set_fontsize(8)
        else:
            ax.text(0.5, 0.5, "No numeric columns to summarize.", ha="center", fontsize=10)
        pdf_pages.savefig(fig)
        plt.close(fig)

        # Data preview page(s)
        preview = df.head(PREVIEW_ROWS)
        preview_rows = [[("" if v is None else str(v)) for v in row] for row in preview.iter_rows()]
        fig2, ax2 = plt.subplots(figsize=(11, 8.5))
        ax2.axis("off")
        ax2.text(0.5, 0.97, f"Data Preview (first {len(preview_rows)} rows)", ha="center", fontsize=13, weight="bold")
        if preview_rows:
            table2 = ax2.table(cellText=preview_rows, colLabels=preview.columns, loc="center", bbox=[0.02, 0.02, 0.96, 0.9])
            table2.auto_set_font_size(False)
            table2.set_fontsize(6)
        pdf_pages.savefig(fig2)
        plt.close(fig2)

    return output_path
